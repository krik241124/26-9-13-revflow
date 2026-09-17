#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
RevFlow
=======

Reverse-engineered catalog workflow automation.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import mimetypes
import os
import re
import sys
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Tuple
from urllib.parse import parse_qsl, urlencode, urlsplit, urlunsplit

import requests
from requests.adapters import HTTPAdapter
from urllib3.util.retry import Retry
from types import SimpleNamespace


# ----------------------------
# Helpers
# ----------------------------
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from project_config import (
    PROJECT_ROOT, load_runtime, market_config, workspace_paths,
    validate_detect, write_extract_context,
)


def load_config() -> dict:
    return load_runtime()


def log(msg: str) -> None:
    print(msg, flush=True)


def format_duration(seconds: float) -> str:
    seconds = max(0, int(round(seconds)))
    if seconds < 60:
        return f"{seconds}s"
    m, s = divmod(seconds, 60)
    if m < 60:
        return f"{m}m{s:02d}s"
    h, m = divmod(m, 60)
    return f"{h}h{m:02d}m"


def safe_name(s: str) -> str:
    return re.sub(r'[^0-9A-Za-z._-]+', '_', str(s)).strip('._') or 'unknown'


def clean_text(v: Any) -> Optional[str]:
    if v is None:
        return None
    s = str(v).strip()
    return s if s else None




def strip_brand_from_title(title: Any, brand: Any) -> Optional[str]:
    """Return an ArkSwift-safe title with the source brand removed.

    Removal is case-insensitive, handles common trademark marks, collapses leftover
    punctuation/whitespace, and keeps the original title only if stripping would
    otherwise produce an empty string.
    """
    original = clean_text(title)
    b = clean_text(brand)
    if not original or not b:
        return original

    # First remove normal token occurrences, including PawHut® / SPORTNOW™.
    pat = re.compile(rf'(?<![\w]){re.escape(b)}(?:[®™©])?(?![\w])', re.IGNORECASE)
    cleaned = pat.sub(' ', original)

    # Distribution titles must contain no source brand. For reasonably distinctive
    # brand strings, also remove a concatenated literal occurrence as a safety net.
    alnum_len = len(re.sub(r'\W+', '', b, flags=re.UNICODE))
    if alnum_len >= 4 and b.casefold() in cleaned.casefold():
        cleaned = re.sub(re.escape(b), ' ', cleaned, flags=re.IGNORECASE)

    cleaned = re.sub(r'\s+', ' ', cleaned)
    cleaned = re.sub(r'^[\s\-–—_|,:;·•/]+', '', cleaned)
    cleaned = re.sub(r'[\s\-–—_|,:;·•/]+$', '', cleaned)
    cleaned = re.sub(r'\s+([,;:])', r'\1', cleaned)
    cleaned = cleaned.strip()
    return cleaned or original


def ensure_brand_free_title(n: Dict[str, Any]) -> bool:
    """Migrate an existing product.json in-place to the brand-free title rule."""
    if not isinstance(n, dict):
        return False
    identity = n.setdefault('identity', {})
    a = n.setdefault('arkswift_fields', {})
    brand = identity.get('brand')
    current = clean_text(a.get('product_title'))
    original = clean_text(a.get('product_title_original')) or current
    if original and not a.get('product_title_original'):
        a['product_title_original'] = original
    cleaned = strip_brand_from_title(original, brand)
    changed = cleaned != current
    if cleaned:
        a['product_title'] = cleaned
    return changed

def first_nonempty(d: Dict[str, Any], *keys: str) -> Any:
    for k in keys:
        if k in d and d[k] not in (None, "", [], {}):
            return d[k]
    return None


def to_number(v: Any) -> Any:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return v
    s = str(v).strip().replace(",", ".")
    try:
        f = float(s)
        return int(f) if f.is_integer() else f
    except Exception:
        return v


def split_codes(v: Any) -> List[str]:
    if v is None:
        return []
    if isinstance(v, list):
        return [str(x).strip() for x in v if str(x).strip()]
    return [x.strip() for x in re.split(r'[|,;\s]+', str(v)) if x.strip()]


def iter_dicts(obj: Any) -> Iterable[Dict[str, Any]]:
    if isinstance(obj, dict):
        yield obj
        for v in obj.values():
            yield from iter_dicts(v)
    elif isinstance(obj, list):
        for x in obj:
            yield from iter_dicts(x)


def find_master(payload: Any) -> Dict[str, Any]:
    """Find the ca_pid_master_all-like object robustly."""
    if isinstance(payload, dict):
        data = payload.get("data")
        if isinstance(data, dict):
            master = data.get("ca_pid_master_all")
            if isinstance(master, dict):
                return master
            if isinstance(master, list) and master and isinstance(master[0], dict):
                return master[0]

    # Fallback: find a dict that looks like the product master record.
    best = None
    best_score = -1
    wanted = {
        "SEGMENT1", "ITEM_NAME_EN", "ITEM_NAME_CN", "BOX_LENGTH",
        "BOX_WIDTH", "BOX_HEIGHT", "GROSS_WEIGHT", "BULLET_POINT1",
        "LONG_DESCRIPTION1", "CATEGORY_NAME3"
    }
    for d in iter_dicts(payload):
        score = sum(1 for k in wanted if k in d)
        if score > best_score:
            best, best_score = d, score
    if best is None or best_score < 3:
        raise ValueError("Could not locate ca_pid_master_all/product master object.")
    return best


# ----------------------------
# cURL -> requests.Session
# ----------------------------

def _decode_cmd_carets(s: str) -> str:
    """Decode Windows CMD caret escaping: ^" -> ", ^& -> &, ^^ -> ^, etc."""
    out = []
    i = 0
    while i < len(s):
        if s[i] == "^" and i + 1 < len(s):
            out.append(s[i + 1])
            i += 2
        else:
            out.append(s[i])
            i += 1
    return "".join(out)


def parse_curl_file(path: Path) -> Tuple[str, Dict[str, str]]:
    """
    Parse Chrome 'Copy as cURL' from:
      - Windows CMD: curl --url ^"https://...?a=1^&b=2^" ^ ...
      - PowerShell/bash-style cURL with quoted URL/headers.

    We keep Cookie as a request header so the authenticated CL browser session is reused.
    """
    import shlex

    raw = path.read_text(encoding="utf-8", errors="ignore").strip()
    if not raw:
        raise ValueError("cURL file is empty.")

    # Remove comments from the template, but keep everything in the actual cURL command.
    raw = "\n".join(
        line for line in raw.splitlines()
        if not line.lstrip().startswith("#")
    ).strip()

    # Join CMD and bash line continuations first.
    raw = re.sub(r"\^\s*\r?\n", " ", raw)
    raw = re.sub(r"\\\s*\r?\n", " ", raw)

    # Chrome on Windows CMD escapes quotes, &, etc. with ^.
    if re.search(r'\^\s*["&|<>^()]|\^"', raw):
        raw = _decode_cmd_carets(raw)

    raw = re.sub(r"\s+", " ", raw).strip()

    # Tokenize after normalization. posix=True correctly handles embedded \" in headers.
    try:
        tokens = shlex.split(raw, posix=True)
    except ValueError as e:
        raise ValueError(f"Could not parse cURL text: {e}") from e

    if not tokens or tokens[0].lower() not in ("curl", "curl.exe"):
        raise ValueError("cURL file does not start with curl/curl.exe.")

    url: Optional[str] = None
    headers: Dict[str, str] = {}

    i = 1
    while i < len(tokens):
        tok = tokens[i]

        if tok in ("--url",):
            if i + 1 >= len(tokens):
                raise ValueError("--url is present but URL value is missing.")
            url = tokens[i + 1]
            i += 2
            continue

        if tok in ("-H", "--header"):
            if i + 1 < len(tokens):
                hv = tokens[i + 1]
                if ":" in hv:
                    name, value = hv.split(":", 1)
                    headers[name.strip()] = value.strip()
                i += 2
                continue

        if tok in ("-b", "--cookie"):
            if i + 1 < len(tokens):
                headers["Cookie"] = tokens[i + 1]
                i += 2
                continue

        # Bash-style Copy as cURL sometimes puts the URL immediately after curl.
        if url is None and tok.startswith(("http://", "https://")):
            url = tok

        i += 1

    if not url:
        raise ValueError(
            "Could not find URL in cURL file. Please use DevTools Network -> "
            "successful GetDetail request -> Copy -> Copy as cURL."
        )

    # Defensive cleanup for any remaining CMD escapes in the URL.
    url = url.replace("^&", "&").replace('^"', '"').strip()

    return url, headers


def make_session(headers: Dict[str, str], cookie_domain: Optional[str] = None) -> requests.Session:
    """
    Build a requests.Session from the copied browser headers.

    Important change in 0.2.2:
    - Browser Cookie header is loaded into Requests' cookie jar instead of being kept
      as a static session header.
    - This lets Set-Cookie responses from CL update the session, matching the browser.
    - When cookie_domain is known, CL auth cookies are scoped to that host so they are
      not sent to the public image CDN.
    """
    from http.cookies import SimpleCookie

    s = requests.Session()
    retry = Retry(
        total=3,
        connect=3,
        read=3,
        status=3,
        backoff_factor=0.5,
        status_forcelist=(429, 500, 502, 503, 504),
        allowed_methods=frozenset(["GET", "HEAD", "POST"]),
        raise_on_status=False,
    )
    s.mount("http://", HTTPAdapter(max_retries=retry))
    s.mount("https://", HTTPAdapter(max_retries=retry))

    bad = {"content-length", "host", ":authority", ":method", ":path", ":scheme", "cookie"}
    cookie_header = None
    for k, v in headers.items():
        if k.lower() == "cookie":
            cookie_header = v
        elif k.lower() not in bad:
            s.headers[k] = v

    if cookie_header:
        cookie = SimpleCookie()
        cookie.load(cookie_header)
        for name, morsel in cookie.items():
            kwargs = {}
            if cookie_domain:
                kwargs["domain"] = cookie_domain
                kwargs["path"] = "/"
            s.cookies.set(name, morsel.value, **kwargs)

    s.headers.setdefault(
        "User-Agent",
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) Chrome/152 Safari/537.36"
    )
    s.headers.setdefault("Accept", "application/json, text/javascript, */*; q=0.01")
    return s


def _query_value(url: str, name: str) -> Optional[str]:
    for k, v in parse_qsl(urlsplit(url).query, keep_blank_values=True):
        if k.lower() == name.lower():
            return v
    return None


def build_sku_url(
    seed_url: str,
    sku: str,
    *,
    line_guid: Optional[str] = None,
    drop_line_guid: bool = False,
) -> str:
    """
    Build a GetDetail URL for a requested SKU.

    Critical CL detail:
    `lineGuid` is not a harmless constant; it identifies a product line record and
    can pin GetDetail to the SKU from which the cURL was copied. Therefore:
      - seed SKU may use the original lineGuid;
      - other SKUs are first tried without lineGuid / with blank lineGuid;
      - if a correct per-SKU MASTER_LINE_GUID is discovered, pass it explicitly.
    """
    parts = urlsplit(seed_url)
    qs = parse_qsl(parts.query, keep_blank_values=True)
    changed_key = False
    saw_line = False
    out = []

    for k, v in qs:
        kl = k.lower()
        if kl == "keyvalue":
            out.append((k, sku))
            changed_key = True
        elif kl == "lineguid":
            saw_line = True
            if drop_line_guid:
                continue
            if line_guid is not None:
                out.append((k, line_guid))
            else:
                out.append((k, v))
        else:
            out.append((k, v))

    if not changed_key:
        raise ValueError(
            "The seed GetDetail URL has no keyValue=... parameter. "
            "Please Copy as cURL from the CL GetDetail XHR request."
        )

    if line_guid is not None and not saw_line and not drop_line_guid:
        out.append(("lineGuid", line_guid))

    return urlunsplit((parts.scheme, parts.netloc, parts.path, urlencode(out), parts.fragment))


def _returned_sku(payload: Any) -> Optional[str]:
    try:
        return clean_text(find_master(payload).get("SEGMENT1"))
    except Exception:
        return None


def _find_line_guid_for_sku(payload: Any, sku: str) -> Optional[str]:
    """Look for a matching SKU/MASTER_LINE_GUID pair anywhere in a GetDetail payload."""
    target = sku.upper()
    for d in iter_dicts(payload):
        seg = clean_text(d.get("SEGMENT1"))
        lg = clean_text(d.get("MASTER_LINE_GUID"))
        if seg and lg and seg.upper() == target:
            return lg
    return None



GET_PAGE_LIST_PATH = "/CA_PID_MASTER/CountryList/GetPageList"


def _country_list_query(sku: str, orgunits_guid: str) -> Dict[str, Any]:
    """
    Reproduce the CountryList/GetPageList queryJson used by the CL UI for an exact SKU lookup.
    The apparently redundant operator fields are preserved because the server-side code may
    expect them to exist even when their paired search value is blank.
    """
    return {
        "ORG_UNITS_GUID": orgunits_guid,
        "LANGUAGE": "",
        "sku": "Equal",
        "SEGMENT1": sku,
        "parent_sku": "Contains",
        "SEGMENT3": "",
        "ORG_UNITS_GUID_NORTH_AMERICA": "",
        "item_name": "Contains",
        "ITEM_NAME_EN": "",
        "item_name_cn": "Contains",
        "ITEM_NAME_CN": "",
        "userSelectCate": "Contains",
        "COPYWRITER_NAME": "",
        "COPYWRITER_NAME1": "Contains",
        "COPYWRITER_SEA_NAME": "",
        "STATUS_FLAG": "",
        "SYNC_FLAG": "",
        "SKU_TYPE": "",
        "FROM_GLOBAL": "",
        "MAIN_LANGUAGE": "",
        "ITEM_STATUS_CODE": "",
        "MAPPING_FLAG": "",
        "SHIP_TIME": "",
        "SHIP_TIME1": "",
        "ARRIVE_TIME": "",
        "ARRIVE_TIME1": "",
        "SKU_FLAG_TYPE": "",
        "PRODUCT_TYPE_NAME": "",
        "SYNC_DATE": "",
        "SYNC_DATE1": "",
        "UPGRADE_SKU": "",
        "CREATE_DATE": "",
        "CREATE_DATE1": "",
        "LAST_UPDATE_DATE": "",
        "LAST_UPDATE_DATE1": "",
        "product_category": "Contains",
        "PRODUCT_CATEGORY": "",
        "product_sub_category": "Contains",
        "PRODUCT_SUB_CATEGORY": "",
        "product_type": "Contains",
        "PRODUCT_TYPE": "",
        "product_category_zhs": "Contains",
        "PRODUCT_CATEGORY_ZHS": "",
        "product_sub_category_zhs": "Contains",
        "PRODUCT_SUB_CATEGORY_ZHS": "",
        "product_type_zhs": "Contains",
        "PRODUCT_TYPE_ZHS": "",
        "scminvnum_op": "Greater",
        "SCM_INVENTORY_NUM": "",
        "SCM_INVENTORY_STATUS": "",
        "scm_first_in_start_time": "",
        "scm_first_in_end_time": "",
        "SUB_RECEIVE_START_TIME": "",
        "SUB_RECEIVE_END_TIME": "",
        "SUB_CONFIRM_START_TIME": "",
        "SUB_CONFIRM_END_TIME": "",
        "IS_EXPLOSIVE_PRODUCT": "",
        "IS_MARKETING_PLAN": "",
        "AOS_HIGHCONVER": "",
    }


def _iter_jsonish_dicts(obj: Any, depth: int = 0) -> Iterable[Dict[str, Any]]:
    """Walk dict/list payloads, including JSON strings nested inside wrapper fields."""
    if depth > 8:
        return
    if isinstance(obj, str):
        t = obj.strip()
        if t and t[:1] in '[{"':
            try:
                decoded = json.loads(t)
            except Exception:
                return
            yield from _iter_jsonish_dicts(decoded, depth + 1)
        return
    if isinstance(obj, dict):
        yield obj
        for v in obj.values():
            yield from _iter_jsonish_dicts(v, depth + 1)
    elif isinstance(obj, list):
        for v in obj:
            yield from _iter_jsonish_dicts(v, depth + 1)


def _resolve_line_guid_from_page_payload(
    payload: Any,
    sku: str,
    language: Optional[str] = None,
) -> Tuple[Optional[str], Dict[str, Any]]:
    """Extract the requested SKU's MASTER_LINE_GUID from GetPageList JSON."""
    target = sku.strip().upper()
    matching_rows: List[Dict[str, Any]] = []
    candidates: List[Dict[str, Any]] = []

    for d in _iter_jsonish_dicts(payload):
        seg = clean_text(first_nonempty(d, "SEGMENT1", "segment1", "SKU", "sku"))
        row_language = clean_text(first_nonempty(
            d,
            "LANGUAGE",
            "COUNTRY_LANGUAGE",
            "LANGUAGE_CODE",
            "LANGUAGE_CD",
        ))
        if not seg or seg.upper() != target:
            continue
        matching_rows.append(d)
        for key in (
            "MASTER_LINE_GUID", "master_line_guid", "MasterLineGuid",
            "LINE_GUID", "lineGuid", "LineGuid",
        ):
            val = clean_text(d.get(key))
            if val:
                candidates.append({
                    "field": key,
                    "value": val,
                    "language": row_language,
                })

    # Exact field wins. De-duplicate while retaining evidence for diagnostics.
    values = []
    for c in candidates:
        if c["value"] not in values:
            values.append(c["value"])

    debug = {
        "requested_sku": sku,
        "matching_row_count": len(matching_rows),
        "line_guid_candidates": candidates,
        "unique_candidate_values": values,
    }
    target_language = (clean_text(language) or "").upper()

    if target_language:
        lang_values = []

        for c in candidates:
            row_lang = (clean_text(c.get("language")) or "").upper()

            if row_lang == target_language and c["value"] not in lang_values:
                lang_values.append(c["value"])

        debug["target_language"] = target_language
        debug["language_candidate_values"] = lang_values

        if len(lang_values) == 1:
            debug["selection_rule"] = f"language={target_language}"
            return lang_values[0], debug
    if len(values) == 1:
        return values[0], debug
    if len(values) > 1:
        # Prefer an explicit MASTER_LINE_GUID if all such values agree.
        explicit = []
        for c in candidates:
            if c["field"].upper() == "MASTER_LINE_GUID" and c["value"] not in explicit:
                explicit.append(c["value"])
        if len(explicit) == 1:
            debug["selection_rule"] = "explicit MASTER_LINE_GUID"
            return explicit[0], debug
    return None, debug


def resolve_line_guid_via_getpagelist(
    session: requests.Session,
    seed_url: str,
    sku: str,
    timeout: int = 30,
) -> Tuple[str, Any, Dict[str, Any]]:
    """
    Resolve SKU -> MASTER_LINE_GUID using the same exact-SKU CountryList/GetPageList
    request the CL UI sends immediately before GetDetail.
    """
    parts = urlsplit(seed_url)
    endpoint = urlunsplit((parts.scheme, parts.netloc, GET_PAGE_LIST_PATH, "", ""))
    orgunits_guid = clean_text(_query_value(seed_url, "country"))
    if not orgunits_guid:
        raise RuntimeError("Seed GetDetail URL has no country=... GUID; cannot call GetPageList.")

    query_json = _country_list_query(sku, orgunits_guid)
    pagination = {
        "rows": 100,
        "page": 1,
        "sidx": "MASTER_HEADER_GUID",
        "sord": "ASC",
        "records": 0,
        "total": 0,
    }
    form = {
        "queryJson": json.dumps(query_json, ensure_ascii=False, separators=(",", ":")),
        "pagination": json.dumps(pagination, ensure_ascii=False, separators=(",", ":")),
    }
    headers = {
        "Accept": "application/json, text/javascript, */*; q=0.01",
        "Content-Type": "application/x-www-form-urlencoded; charset=UTF-8",
        "X-Requested-With": "XMLHttpRequest",
        "Origin": f"{parts.scheme}://{parts.netloc}",
    }

    r = session.post(endpoint, data=form, headers=headers, timeout=timeout, allow_redirects=True)
    if looks_logged_out(r):
        raise RuntimeError(
            f"CL session appears expired/logged out while resolving lineGuid "
            f"(HTTP {r.status_code}, final_url={r.url})."
        )
    r.raise_for_status()
    try:
        payload = r.json()
    except Exception:
        body = re.sub(r"\s+", " ", r.text[:700]).strip()
        raise RuntimeError(
            f"GetPageList did not return JSON (HTTP {r.status_code}, "
            f"content-type={r.headers.get('content-type','')}, body_start={body!r})."
        )

    desired_language = clean_text(_query_value(seed_url, "language"))

    guid, debug = _resolve_line_guid_from_page_payload(
        payload,
        sku,
        desired_language,
    )
    debug.update({
        "endpoint": endpoint,
        "http": r.status_code,
        "payload_shape": _payload_shape(payload) if "_payload_shape" in globals() else type(payload).__name__,
        "queryJson": query_json,
        "pagination": pagination,
    })
    if not guid:
        raise LineGuidResolutionError(
            "GetPageList returned JSON but no unique MASTER_LINE_GUID was found for "
            f"{sku}. See raw_getpagelist.json and lineguid_resolution.json.",
            payload=payload,
            debug=debug,
        )
    debug["resolved_line_guid"] = guid
    return guid, payload, debug


class LineGuidResolutionError(RuntimeError):
    def __init__(self, message: str, payload: Any = None, debug: Optional[Dict[str, Any]] = None):
        super().__init__(message)
        self.payload = payload
        self.debug = debug or {}


class GetDetailResolutionError(RuntimeError):
    def __init__(self, message: str, attempts: List[Dict[str, Any]]):
        super().__init__(message)
        self.attempts = attempts


def fetch_getdetail_for_sku(
    session: requests.Session,
    seed_url: str,
    sku: str,
    timeout: int = 30,
    resolved_line_guid: Optional[str] = None,
) -> Tuple[Dict[str, Any], str, List[Dict[str, Any]]]:
    """
    Fetch GetDetail and GUARANTEE that returned SEGMENT1 == requested SKU.

    CL requires a SKU-specific lineGuid. In 0.4 the normal batch path resolves it via
    CountryList/GetPageList first. This function validates the returned SKU and refuses
    silent seed-SKU contamination.
    """
    seed_sku = clean_text(_query_value(seed_url, "keyValue"))
    seed_line_guid = clean_text(_query_value(seed_url, "lineGuid"))
    sku_u = sku.upper()
    attempts: List[Dict[str, Any]] = []

    variants: List[Tuple[str, str]] = []
    if resolved_line_guid:
        variants.append(("getpagelist-lineGuid", build_sku_url(seed_url, sku, line_guid=resolved_line_guid)))
    elif seed_sku and seed_sku.upper() == sku_u:
        variants.append(("seed-lineGuid", build_sku_url(seed_url, sku)))
    else:
        # Diagnostic fallbacks only. The normal 0.4 batch path resolves the real
        # per-SKU lineGuid through CountryList/GetPageList before calling here.
        variants.append(("drop-lineGuid", build_sku_url(seed_url, sku, drop_line_guid=True)))
        variants.append(("blank-lineGuid", build_sku_url(seed_url, sku, line_guid="")))

    tried_urls = set()

    def attempt(label: str, url: str) -> Optional[Tuple[Dict[str, Any], str]]:
        if url in tried_urls:
            return None
        tried_urls.add(url)
        rec: Dict[str, Any] = {"mode": label, "url": url}
        try:
            r = session.get(url, timeout=timeout, allow_redirects=True)
            rec["http"] = r.status_code
            rec["final_url"] = r.url
            if looks_logged_out(r):
                rec["result"] = "logged_out"
                attempts.append(rec)
                raise RuntimeError("CL session appears expired/logged out while fetching GetDetail.")
            r.raise_for_status()
            try:
                payload = r.json()
            except Exception:
                rec["result"] = "non_json"
                rec["body_start"] = re.sub(r"\s+", " ", r.text[:500]).strip()
                attempts.append(rec)
                return None

            rec["code"] = payload.get("code") if isinstance(payload, dict) else None
            returned = _returned_sku(payload)
            try:
                master = find_master(payload)
                returned_line_guid = clean_text(master.get("MASTER_LINE_GUID"))
            except Exception:
                returned_line_guid = None

            rec["returned_sku"] = returned
            rec["returned_line_guid"] = returned_line_guid

            sku_match = bool(returned and returned.upper() == sku_u)
            is_seed_sku = bool(seed_sku and seed_sku.upper() == sku_u)

            # For non-seed SKUs, a returned SEGMENT1 alone is not enough. We observed
            # CL can mix the requested SEGMENT1 with localized/detail fields pinned by
            # the original seed lineGuid. Require a real, non-seed MASTER_LINE_GUID.
            line_ok = is_seed_sku or bool(
                returned_line_guid and
                (not seed_line_guid or returned_line_guid.lower() != seed_line_guid.lower())
            )

            if sku_match and line_ok:
                rec["result"] = "match"
            elif sku_match:
                rec["result"] = "sku_match_but_seed_lineGuid_contamination"
            else:
                rec["result"] = "mismatch"
            attempts.append(rec)

            if isinstance(payload, dict) and payload.get("code") not in (None, 200, "200"):
                return None

            if sku_match and line_ok:
                return payload, r.url

            # Sometimes a payload can expose itemList/sibling lines. If the requested
            # SKU appears there, immediately retry with its real MASTER_LINE_GUID.
            found_line = _find_line_guid_for_sku(payload, sku)
            if found_line:
                exact_url = build_sku_url(seed_url, sku, line_guid=found_line)
                exact = attempt("discovered-lineGuid", exact_url)
                if exact:
                    return exact
            return None
        except RuntimeError:
            raise
        except Exception as e:
            rec.setdefault("result", "request_error")
            rec["error"] = str(e)
            if rec not in attempts:
                attempts.append(rec)
            return None

    for label, url in variants:
        result = attempt(label, url)
        if result:
            payload, final_url = result
            return payload, final_url, attempts

    returned_values = [a.get("returned_sku") for a in attempts if a.get("returned_sku")]
    returned_msg = ", ".join(dict.fromkeys(returned_values)) if returned_values else "none"
    raise GetDetailResolutionError(
        "GetDetail did not resolve the requested SKU. "
        f"requested={sku}; returned={returned_msg}. "
        "The copied GetDetail lineGuid is SKU-specific. "
        "0.3 refused to save another SKU's data. "
        "0.4 normally resolves MASTER_LINE_GUID through CountryList/GetPageList first; "
        "inspect raw_getpagelist.json/lineguid_resolution.json if resolution failed.",
        attempts,
    )


def looks_logged_out(resp: requests.Response) -> bool:
    """
    Detect a real login/session-expiry response without false-positives on CL JSON.

    Important: CL photo rows contain a field named LAST_UPDATE_LOGIN. Some ASP.NET
    endpoints may still label JSON as text/html, so searching the raw body for the
    substring "login" is NOT safe.
    """
    if resp.status_code in (401, 403):
        return True

    # If the body is valid JSON (even when mislabeled as text/html), it is not a login page.
    try:
        resp.json()
        return False
    except Exception:
        pass

    final_url = (resp.url or "").lower()
    if re.search(r"/(?:login|signin|sign-in)(?:[/?#]|$)", final_url):
        return True

    ctype = (resp.headers.get("content-type") or "").lower()
    text = resp.text[:12000]
    low = text.lower()

    # Strong HTML login-page signals only. Do not match generic data-field names.
    if "text/html" in ctype or "<html" in low or "<!doctype html" in low:
        if re.search(r"<input[^>]+type=[\"']?password", low):
            return True
        if re.search(r"<form[^>]+(?:login|signin|sign-in)", low):
            return True
        if re.search(r"<title[^>]*>[^<]*(?:login|sign in|signin|登录)[^<]*</title>", low):
            return True
        if ("asp.net_sessionid" in low and "password" in low and "username" in low):
            return True
    return False


# ----------------------------
# Input
# ----------------------------

def load_skus(path: Path, sheet: Optional[str] = None) -> List[str]:
    ext = path.suffix.lower()

    if ext in (".txt", ".list"):
        vals = [x.strip() for x in path.read_text(encoding="utf-8-sig").splitlines()]
        return list(dict.fromkeys(x for x in vals if x and not x.startswith("#")))

    if ext == ".csv":
        with path.open("r", encoding="utf-8-sig", newline="") as f:
            rows = list(csv.reader(f))
        if not rows:
            return []
        header = [str(x).strip().lower() for x in rows[0]]
        idx = header.index("sku") if "sku" in header else 0
        vals = [r[idx].strip() for r in rows[1:] if len(r) > idx and r[idx].strip()]
        return list(dict.fromkeys(vals))

    if ext in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            raise RuntimeError("Excel input requires openpyxl: pip install openpyxl")
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[sheet] if sheet else wb.active
        rows = ws.iter_rows(values_only=True)
        first = next(rows, None)
        if not first:
            return []
        header = [str(x).strip().lower() if x is not None else "" for x in first]
        idx = header.index("sku") if "sku" in header else 0
        vals = []
        for r in rows:
            if len(r) > idx and r[idx] is not None:
                v = str(r[idx]).strip()
                if v:
                    vals.append(v)
        return list(dict.fromkeys(vals))

    raise ValueError(f"Unsupported input type: {ext}. Use .txt, .csv, .xlsx or .xlsm.")


# ----------------------------
# Normalization for ArkSwift/RPA
# ----------------------------

def normalize_product(payload: Dict[str, Any], requested_sku: str, country: str) -> Dict[str, Any]:
    m = find_master(payload)

    bullets = [
        clean_text(m.get(f"BULLET_POINT{i}"))
        for i in range(1, 6)
    ]
    bullets = [x for x in bullets if x]

    eans = split_codes(first_nonempty(m, "EAN_CODE1", "EAN_CODE", "UPC", "EAN"))

    title = first_nonempty(
        m,
        "AOSOM_TITLE",
        "AMAZON_ITEM_NAME",
        "AMAZON_TITLE",
        "ITEM_NAME_EN",
        "ITEM_NAME_CN",
    )
    brand = clean_text(first_nonempty(m, "BRAND", "BRAND1"))
    original_title = clean_text(title)
    arkswift_title = strip_brand_from_title(original_title, brand)

    direct_images = []
    for k in ("PIC_URL", "PICTURE_URL", "URL_IMAGE"):
        v = clean_text(m.get(k))
        if v and v.startswith(("http://", "https://")) and v not in direct_images:
            direct_images.append(v)

    result = {
        "schema_version": "cl-to-arkswift-mvp-0.5",
        "source": {
            "system": "CL",
            "country": country,
            "requested_sku": requested_sku,
            "cl_ids": {
                "standitem_guid": clean_text(m.get("STANDAND_ITEM_GUID")),
                "orgunits_guid": clean_text(m.get("ORG_UNITS_GUID")),
                "orglang_guid": clean_text(m.get("ORG_LANGUAGE_GUID")),
                "masterheader_guid": clean_text(m.get("MASTER_HEADER_GUID")),
            },
        },
        "identity": {
            "seller_sku": clean_text(first_nonempty(m, "SEGMENT1")) or requested_sku,
            "item_name_cn": clean_text(m.get("ITEM_NAME_CN")),
            "item_name_en": clean_text(m.get("ITEM_NAME_EN")),
            "brand": brand,
            "parent_sku": clean_text(first_nonempty(m, "PARENT_SKU", "PARENT_SEGMENT1")),
        },
        "category_source": {
            "level1": clean_text(m.get("CATEGORY_NAME1")),
            "level2": clean_text(m.get("CATEGORY_NAME2")),
            "level3": clean_text(m.get("CATEGORY_NAME3")),
            "product_type": clean_text(first_nonempty(m, "PRODUCT_TYPE2", "PRODUCT_TYPE")),
            "product_type_name": clean_text(m.get("PRODUCT_TYPE_NAME")),
        },
        "arkswift_fields": {
            "product_title": arkswift_title,
            "product_title_original": original_title,
            "seller_sku": clean_text(first_nonempty(m, "SEGMENT1")) or requested_sku,
            "upc_ean_candidates": eans,
            "country_of_origin": clean_text(m.get("COUNTRY_CD")),
            "color_source": clean_text(first_nonempty(m, "COLOR1", "COLOR", "COLOR_MAP")),
            "color_map_source": clean_text(m.get("COLOR_MAP")),
            "material_source": clean_text(first_nonempty(m, "MATERIAL")),
            # Do not guess business-only fields.
            "is_white_label": None,
            "is_original": None,
            "assembled_dimensions_cm": {
                "length": to_number(m.get("LENGTH")),
                "width": to_number(m.get("WIDTH")),
                "height": to_number(m.get("HEIGHT")),
            },
            "net_weight_kg": to_number(m.get("NET_WEIGHT")),
            "package": {
                "quantity": to_number(first_nonempty(m, "QUANTITY", "BOXES_QUANTITY")),
                "length_cm": to_number(m.get("BOX_LENGTH")),
                "width_cm": to_number(m.get("BOX_WIDTH")),
                "height_cm": to_number(m.get("BOX_HEIGHT")),
                "gross_weight_kg": to_number(m.get("GROSS_WEIGHT")),
            },
            "features": bullets,
            "description": clean_text(first_nonempty(m, "LONG_DESCRIPTION1", "AOSOM_LONG_DESCRIPTION1")),
            "specifications": clean_text(first_nonempty(m, "SPECIFICATION1", "AOSOM_SPECIFICATION1")),
            "package_includes": clean_text(first_nonempty(m, "PACKAGE_INCLUDES1", "AOSOM_PACKAGE_INCLUDES1")),
            "electric_flag": clean_text(m.get("PRODUCT_ELECTRIC_FLAG")),
            "manual_url": clean_text(first_nonempty(m, "PRODUCT_URL", "URL_INSTRUCTION")),
        },
        "image_source": {
            "direct_urls_from_getdetail": direct_images,
            "logical_photo_list": [],
            "gpsr_urls": [],
        },
        "raw_hints": {
            "item_specification_cn": clean_text(m.get("ITEM_SPECIFICATION_CN")),
            "item_attribute1": clean_text(m.get("ITEM_ATTRIBUTE1")),
            "item_attribute2": clean_text(m.get("ITEM_ATTRIBUTE2")),
            "differences": clean_text(m.get("Differences")),
        },
        "validation": {
            "missing_required": [],
            "warnings": [],
        },
    }

    required = {
        "product_title": result["arkswift_fields"]["product_title"],
        "seller_sku": result["arkswift_fields"]["seller_sku"],
        "material_source": result["arkswift_fields"]["material_source"],
        "color_source": result["arkswift_fields"]["color_source"],
        "length_cm": result["arkswift_fields"]["assembled_dimensions_cm"]["length"],
        "width_cm": result["arkswift_fields"]["assembled_dimensions_cm"]["width"],
        "height_cm": result["arkswift_fields"]["assembled_dimensions_cm"]["height"],
        "net_weight_kg": result["arkswift_fields"]["net_weight_kg"],
        "box_length_cm": result["arkswift_fields"]["package"]["length_cm"],
        "box_width_cm": result["arkswift_fields"]["package"]["width_cm"],
        "box_height_cm": result["arkswift_fields"]["package"]["height_cm"],
        "gross_weight_kg": result["arkswift_fields"]["package"]["gross_weight_kg"],
    }
    result["validation"]["missing_required"] = [k for k, v in required.items() if v in (None, "", [])]

    if not eans:
        result["validation"]["warnings"].append("No EAN/UPC candidate found.")
    if not bullets:
        result["validation"]["warnings"].append("No bullet points found.")
    if not result["image_source"]["direct_urls_from_getdetail"]:
        result["validation"]["warnings"].append("No direct image URL found in GetDetail.")

    return result


# ----------------------------
# Images
# ----------------------------

PHOTO_PAGE_PATH = "/CA_PID_MASTER/OrgPidFenceInfo/Photo"
PHOTO_LIST_PATH = "/CA_PID_MASTER/OrgPidFenceInfo/GetListForPhoto"
DEFAULT_IMAGE_CDN = "https://img-cl.aosomcdn.com"


def ext_from_content_type(ctype: str, fallback_url: str = "") -> str:
    ctype = (ctype or "").split(";")[0].strip().lower()
    mapping = {
        "image/jpeg": ".jpg",
        "image/jpg": ".jpg",
        "image/png": ".png",
        "image/gif": ".gif",
        "image/webp": ".webp",
    }
    if ctype in mapping:
        return mapping[ctype]
    suffix = Path(urlsplit(fallback_url).path).suffix.lower()
    if suffix in (".jpg", ".jpeg", ".png", ".gif", ".webp"):
        return ".jpg" if suffix == ".jpeg" else suffix
    return ".jpg"


def download_image(session: requests.Session, url: str, dest_stem: Path, timeout: int = 20) -> Optional[Path]:
    try:
        r = session.get(url, timeout=timeout, stream=True, allow_redirects=True)
        if r.status_code != 200:
            return None
        ctype = (r.headers.get("content-type") or "").lower()
        if not ctype.startswith("image/"):
            return None
        ext = ext_from_content_type(ctype, url)
        dest = dest_stem.with_suffix(ext)
        with dest.open("wb") as f:
            for chunk in r.iter_content(1024 * 128):
                if chunk:
                    f.write(chunk)
        return dest
    except Exception:
        return None


def _jsonish_decode(value: Any, max_rounds: int = 4) -> Any:
    """
    Decode JSON that may itself be wrapped as a JSON string one or more times.
    Some CL/Learun endpoints return e.g. {"data": "[{...}]"}.
    """
    current = value
    for _ in range(max_rounds):
        if not isinstance(current, str):
            break
        s = current.strip()
        if not s:
            break
        # Fast guard: only attempt strings that plausibly contain JSON.
        if s[0] not in "[{\"":
            break
        try:
            nxt = json.loads(s)
        except Exception:
            break
        if nxt == current:
            break
        current = nxt
    return current


def _photo_rows_from_payload(payload: Any) -> List[Dict[str, Any]]:
    """
    Find CL photo rows robustly across:
      - raw array
      - {"data": [...]}
      - {"data": "[...]"}  (JSON string wrapper)
      - nested/double-encoded wrappers
      - arbitrary dict/list containers

    Candidate arrays are scored by PIC_URL / PIC_URL_SECRET / NAME presence.
    """
    candidates: List[Tuple[int, List[Dict[str, Any]]]] = []
    visited_ids = set()

    def walk(o: Any, depth: int = 0) -> None:
        if depth > 10:
            return

        o2 = _jsonish_decode(o)
        if o2 is not o:
            walk(o2, depth + 1)
            return

        if isinstance(o, (dict, list)):
            oid = id(o)
            if oid in visited_ids:
                return
            visited_ids.add(oid)

        if isinstance(o, list):
            rows = [x for x in o if isinstance(x, dict)]
            if rows:
                pic = sum(1 for x in rows if x.get("PIC_URL") not in (None, ""))
                secret = sum(1 for x in rows if x.get("PIC_URL_SECRET") not in (None, ""))
                named = sum(1 for x in rows if x.get("NAME") not in (None, ""))
                # Even arrays with null PIC_URL rows are valid CL photoObj data.
                if any(("PIC_URL" in x or "PIC_URL_SECRET" in x or "NAME" in x) for x in rows):
                    score = pic * 10 + secret * 5 + named + len(rows)
                    candidates.append((score, rows))
            for x in o:
                walk(x, depth + 1)

        elif isinstance(o, dict):
            # A single photo row can occur by itself.
            if any(k in o for k in ("PIC_URL", "PIC_URL_SECRET", "NAME")):
                score = (10 if o.get("PIC_URL") else 0) + (5 if o.get("PIC_URL_SECRET") else 0) + 1
                candidates.append((score, [o]))
            for v in o.values():
                walk(v, depth + 1)

        elif isinstance(o, str):
            # Try URL-decoded or HTML-escaped JSON as a last resort.
            from urllib.parse import unquote
            import html
            for transformed in (html.unescape(o), unquote(o)):
                if transformed != o:
                    parsed = _jsonish_decode(transformed)
                    if parsed is not transformed:
                        walk(parsed, depth + 1)

    walk(payload)
    if not candidates:
        return []
    candidates.sort(key=lambda x: (x[0], len(x[1])), reverse=True)
    return candidates[0][1]


def _payload_shape(payload: Any) -> str:
    """Compact diagnostic description that is safe to print in an error."""
    if isinstance(payload, list):
        return f"list(len={len(payload)})"
    if isinstance(payload, dict):
        parts = []
        for k, v in list(payload.items())[:12]:
            if isinstance(v, list):
                parts.append(f"{k}=list({len(v)})")
            elif isinstance(v, dict):
                parts.append(f"{k}=dict({len(v)})")
            elif isinstance(v, str):
                parts.append(f"{k}=str({len(v)})")
            else:
                parts.append(f"{k}={type(v).__name__}")
        return "dict(" + ", ".join(parts) + ")"
    if isinstance(payload, str):
        return f"str(len={len(payload)})"
    return type(payload).__name__


def _photo_page_url(seed_url: str, normalized: Dict[str, Any]) -> str:
    ids = normalized.get("source", {}).get("cl_ids", {})
    parts = urlsplit(seed_url)
    endpoint = urlunsplit((parts.scheme, parts.netloc, PHOTO_PAGE_PATH, "", ""))
    params = {
        "masterheaderValue": ids["masterheader_guid"],
        "orgunitsValue": ids["orgunits_guid"],
        "orglangValue": ids["orglang_guid"],
        "standitemValue": ids["standitem_guid"],
    }
    return requests.Request("GET", endpoint, params=params).prepare().url


def fetch_photo_list(
    session: requests.Session,
    seed_url: str,
    normalized: Dict[str, Any],
    timeout: int = 30,
) -> Tuple[Any, List[Dict[str, Any]], str, Dict[str, Any]]:
    """
    Mirror the real browser sequence:

      1) GET /OrgPidFenceInfo/Photo?...          (prime page/session state)
      2) GET /OrgPidFenceInfo/GetListForPhoto?... (XHR)

    0.2/0.2.1 skipped step (1). If CL initializes server-side state or refreshes
    cookies while opening the Photo iframe, the direct XHR can return an empty payload.
    """
    ids = normalized.get("source", {}).get("cl_ids", {})
    missing = [k for k in ("standitem_guid", "orgunits_guid", "orglang_guid", "masterheader_guid") if not ids.get(k)]
    if missing:
        raise RuntimeError("GetDetail is missing CL photo-list GUID(s): " + ", ".join(missing))

    # Prime the Photo document exactly as the browser does before its image-list XHR.
    page_url = _photo_page_url(seed_url, normalized)
    page_resp = session.get(page_url, timeout=timeout, allow_redirects=True)
    if looks_logged_out(page_resp):
        raise RuntimeError(
            f"CL session appears expired/logged out while opening Photo page "
            f"(HTTP {page_resp.status_code}, final_url={page_resp.url})."
        )
    page_resp.raise_for_status()

    parts = urlsplit(seed_url)
    endpoint = urlunsplit((parts.scheme, parts.netloc, PHOTO_LIST_PATH, "", ""))
    strgloble = json.dumps({
        "standitemValue": ids["standitem_guid"],
        "orgunitsValue": ids["orgunits_guid"],
        "orglangValue": ids["orglang_guid"],
        "masterheaderValue": ids["masterheader_guid"],
    }, ensure_ascii=False, separators=(",", ":"))

    params = {
        "shelfitemValue": "",
        "strgloblePid": strgloble,
        "_": str(int(time.time() * 1000)),
    }

    # Match browser XHR semantics more closely.
    xhr_headers = {
        "Accept": "application/json, text/javascript, */*; q=0.01",
        "X-Requested-With": "XMLHttpRequest",
        # The copied browser request showed Home/Index as Referer. Keep the existing
        # copied Referer if present instead of forcing the Photo URL.
    }

    r = session.get(endpoint, params=params, headers=xhr_headers, timeout=timeout, allow_redirects=True)
    if looks_logged_out(r):
        raise RuntimeError(
            f"CL session appears expired/logged out while fetching photo list "
            f"(HTTP {r.status_code}, final_url={r.url})."
        )
    r.raise_for_status()

    try:
        payload = r.json()
    except Exception as e:
        ctype = r.headers.get("content-type", "")
        snippet = re.sub(r"\s+", " ", r.text[:500]).strip()
        raise RuntimeError(
            "GetListForPhoto did not return JSON. "
            f"HTTP={r.status_code}; content-type={ctype!r}; final_url={r.url}; "
            f"body_start={snippet!r}"
        ) from e

    rows = _photo_rows_from_payload(payload)
    debug = {
        "photo_page_url": page_resp.url,
        "photo_page_http": page_resp.status_code,
        "photo_page_content_type": page_resp.headers.get("content-type"),
        "photo_list_url": r.url,
        "photo_list_http": r.status_code,
        "photo_list_content_type": r.headers.get("content-type"),
        "payload_shape": _payload_shape(payload),
        "response_body_start": re.sub(r"\s+", " ", r.text[:1200]).strip(),
        "cookies_after_photo_page": sorted([c.name for c in session.cookies]),
    }
    return payload, rows, r.url, debug


def _rank_value(row: Dict[str, Any], fallback: int) -> float:
    try:
        return float(row.get("RANK"))
    except Exception:
        return float(fallback)


def _logical_name(row: Dict[str, Any], index: int) -> str:
    name = clean_text(row.get("NAME"))
    if name:
        return name
    path = clean_text(row.get("PIC_URL")) or f"unnamed_{index}"
    # Collapse localized and generic variants to the same path key where possible.
    path = re.sub(r"/([A-Z]{2})/", "/", path, count=1)
    return path


def normalize_photo_rows(rows: List[Dict[str, Any]], sku: str, country: str) -> Dict[str, Any]:
    """
    Build RPA-friendly logical image slots.

    CL commonly returns two records for one logical NAME: e.g. FR/1-1.jpg and 1-1.jpg.
    We keep both candidates but prefer /FR/ first. GPSR and /nobody/ are separated.
    """
    sku_upper = sku.upper()
    country_upper = country.upper()
    usable: List[Tuple[int, Dict[str, Any]]] = []
    gpsr: List[Dict[str, Any]] = []
    ignored: List[Dict[str, Any]] = []

    for i, row in enumerate(rows):
        path = clean_text(row.get("PIC_URL"))
        secret = clean_text(row.get("PIC_URL_SECRET"))
        segment = (clean_text(row.get("SEGMENT1")) or "").upper()
        if not path:
            continue
        if path.upper().startswith("GPSR/"):
            gpsr.append(row)
            continue
        if "/NOBODY/" in path.upper():
            ignored.append(row)
            continue
        if segment and segment != sku_upper:
            continue
        if not secret:
            ignored.append(row)
            continue
        usable.append((i, row))

    groups: Dict[str, List[Tuple[int, Dict[str, Any]]]] = {}
    for i, row in usable:
        groups.setdefault(_logical_name(row, i), []).append((i, row))

    logical: List[Dict[str, Any]] = []
    for name, group in groups.items():
        def cand_sort(pair: Tuple[int, Dict[str, Any]]) -> Tuple[int, float, int]:
            i, row = pair
            path = clean_text(row.get("PIC_URL")) or ""
            localized = f"/{country_upper}/" in path.upper()
            return (0 if localized else 1, _rank_value(row, i), i)

        group_sorted = sorted(group, key=cand_sort)
        min_rank = min(_rank_value(row, i) for i, row in group)
        candidates = []
        for i, row in group_sorted:
            path = clean_text(row.get("PIC_URL")) or ""
            candidates.append({
                "path": path,
                "secret": clean_text(row.get("PIC_URL_SECRET")),
                "rank": row.get("RANK"),
                "localized": f"/{country_upper}/" in path.upper(),
                "fence_guid": clean_text(row.get("FENCE_GUID")),
            })
        logical.append({
            "name": name,
            "rank": min_rank,
            "candidates": candidates,
        })

    logical.sort(key=lambda x: (x["rank"], x["name"]))
    return {
        "logical": logical,
        "gpsr": gpsr,
        "ignored_count": len(ignored),
        "raw_row_count": len(rows),
        "usable_row_count": len(usable),
    }


def _fetch_slot_candidate_bytes(
    slot: Dict[str, Any],
    image_cdn: str,
    timeout: int,
    base_headers: Optional[Dict[str, str]] = None,
) -> Dict[str, Any]:
    """Worker used only for CDN image traffic. CL API calls remain serial."""
    attempts = []
    headers = {}
    if base_headers:
        ua = base_headers.get("User-Agent") or base_headers.get("user-agent")
        if ua:
            headers["User-Agent"] = ua
    headers.setdefault("Accept", "image/avif,image/webp,image/apng,image/svg+xml,image/*,*/*;q=0.8")

    for cand in slot.get("candidates", []):
        secret = cand.get("secret")
        if not secret:
            continue
        url = image_cdn.rstrip("/") + "/" + str(secret).lstrip("/")
        attempt = {
            "path": cand.get("path"),
            "url": url,
            "localized": bool(cand.get("localized")),
        }
        try:
            r = requests.get(url, headers=headers, timeout=timeout, allow_redirects=True)
            attempt["http"] = r.status_code
            ctype = (r.headers.get("content-type") or "").lower()
            attempt["content_type"] = ctype
            attempts.append(attempt)
            if r.status_code == 200 and ctype.startswith("image/") and r.content:
                return {
                    "ok": True,
                    "content": r.content,
                    "content_type": ctype,
                    "url": url,
                    "candidate": cand,
                    "attempts": attempts,
                }
        except Exception as e:
            attempt["error"] = str(e)
            attempts.append(attempt)
    return {"ok": False, "attempts": attempts}


def _write_manifest(path: Path, manifest: Dict[str, Any]) -> None:
    path.write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")


def collect_exact_images(
    session: requests.Session,
    seed_url: str,
    normalized: Dict[str, Any],
    sku_dir: Path,
    max_images: int,
    image_cdn: str,
    download_gpsr: bool,
    timeout: int,
    image_workers: int = 4,
) -> Dict[str, Any]:
    """Fetch exact CL photo metadata serially, download CDN images concurrently, stop at target."""
    out_dir = sku_dir / "images"
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest_path = out_dir / "manifest.json"

    # 0 is intentionally treated as ArkSwift's normal 12-image target for backward
    # compatibility with old command lines that used --max-images 0.
    unlimited = max_images < 0
    target = None if unlimited else (12 if max_images == 0 else max(1, max_images))
    target_label = "ALL" if unlimited else str(target)

    # Reuse photo-list metadata from disk on resume; this avoids another CL API call.
    raw_photo_path = sku_dir / "raw_photolist.json"
    debug_path = sku_dir / "photolist_debug.json"
    raw_payload = None
    rows: List[Dict[str, Any]] = []
    request_url = "reused-local-raw_photolist.json"
    fetch_debug: Dict[str, Any] = {"reused_local": False}
    if raw_photo_path.exists():
        try:
            raw_payload = json.loads(raw_photo_path.read_text(encoding="utf-8"))
            rows = _photo_rows_from_payload(raw_payload)
            if rows:
                fetch_debug = {"reused_local": True, "payload_shape": _payload_shape(raw_payload)}
                log("  [PHOTO] reused local raw_photolist.json; no CL photo-list API call")
        except Exception:
            raw_payload = None
            rows = []

    if not rows:
        raw_payload, rows, request_url, fetch_debug = fetch_photo_list(
            session, seed_url, normalized, timeout=timeout
        )
        raw_photo_path.write_text(
            json.dumps(raw_payload, ensure_ascii=False, indent=2), encoding="utf-8"
        )
        debug_path.write_text(
            json.dumps(fetch_debug, ensure_ascii=False, indent=2), encoding="utf-8"
        )

    if not rows:
        raise RuntimeError(
            "GetListForPhoto returned JSON but no photo rows were found. "
            f"payload={fetch_debug.get('payload_shape')}. "
            "See raw_photolist.json and photolist_debug.json."
        )

    sku = normalized["identity"]["seller_sku"] or normalized["source"]["requested_sku"]
    country = normalized["source"]["country"]
    info = normalize_photo_rows(rows, sku, country)
    normalized["image_source"]["logical_photo_list"] = info["logical"]
    normalized["image_source"]["gpsr_urls"] = [
        clean_text(x.get("PIC_URL")) for x in info["gpsr"] if clean_text(x.get("PIC_URL"))
    ]

    # Resume an interrupted 0.5 image run. A pre-0.5 manifest has no `complete` flag;
    # because old versions only wrote it at the end, such a manifest is already complete.
    previous: Dict[str, Any] = {}
    if manifest_path.exists():
        try:
            previous = json.loads(manifest_path.read_text(encoding="utf-8"))
        except Exception:
            previous = {}

    if previous and "complete" not in previous:
        previous["complete"] = True

    manifest: Dict[str, Any] = {
        "mode": "exact",
        "version": "0.5",
        "complete": False,
        "target_images": target_label,
        "photo_list_endpoint": PHOTO_LIST_PATH,
        "photo_list_request_url": request_url,
        "raw_row_count": info["raw_row_count"],
        "usable_row_count": info["usable_row_count"],
        "logical_slot_count": len(info["logical"]),
        "ignored_count": info["ignored_count"],
        "downloaded": [],
        "failed_slots": [],
        "duplicate_content": [],
        "processed_slots": [],
        "arkswift_upload_candidates": [],
    }

    # Carry forward only checkpoint records whose files still exist.
    seen_hashes = set()
    processed_names = set()
    if previous:
        for item in previous.get("downloaded", []):
            f = out_dir / str(item.get("file", ""))
            if f.is_file():
                manifest["downloaded"].append(item)
                h = clean_text(item.get("sha256"))
                if h:
                    seen_hashes.add(h)
                name = clean_text(item.get("name"))
                if name:
                    processed_names.add(name)
        for item in previous.get("duplicate_content", []):
            manifest["duplicate_content"].append(item)
            name = clean_text(item.get("name"))
            if name:
                processed_names.add(name)
        for item in previous.get("failed_slots", []):
            manifest["failed_slots"].append(item)
            name = clean_text(item.get("name"))
            if name:
                processed_names.add(name)
        processed_names.update(str(x) for x in previous.get("processed_slots", []) if x)

    # If enough images already exist, finish immediately.
    if target is not None and len(manifest["downloaded"]) >= target:
        manifest["downloaded"] = manifest["downloaded"][:target]
        manifest["complete"] = True
        manifest["processed_slots"] = sorted(processed_names)
        manifest["arkswift_upload_candidates"] = [x["file"] for x in manifest["downloaded"][:12]]
        normalized["image_source"]["downloaded_files"] = [x["file"] for x in manifest["downloaded"]]
        normalized["image_source"]["arkswift_upload_candidates"] = manifest["arkswift_upload_candidates"]
        _write_manifest(manifest_path, manifest)
        log(f"  [IMAGES] resume: already have {len(manifest['downloaded'])}/{target}; no CDN work")
        return manifest

    slots = [x for x in info["logical"] if x.get("name") not in processed_names]
    total_slots = len(info["logical"])
    already_processed = total_slots - len(slots)
    workers = max(1, min(int(image_workers or 1), 8))
    img_started = time.time()
    checked = already_processed

    log(
        f"  [IMAGES] logical slots={total_slots}; target={target_label}; "
        f"workers={workers}; resumed={already_processed}; starting CDN downloads..."
    )
    _write_manifest(manifest_path, manifest)

    # Process rank-ordered slots in small concurrent batches. We only schedule the
    # next batch if the target has not yet been reached, so 12 images stops the run.
    cursor = 0
    while cursor < len(slots):
        if target is not None and len(manifest["downloaded"]) >= target:
            break
        batch_size = workers
        if target is not None:
            batch_size = min(workers, max(1, target - len(manifest["downloaded"])))
        batch = slots[cursor:cursor + batch_size]
        cursor += len(batch)
        results: Dict[str, Dict[str, Any]] = {}

        with ThreadPoolExecutor(max_workers=workers, thread_name_prefix="cl-img") as pool:
            futures = {
                pool.submit(
                    _fetch_slot_candidate_bytes,
                    slot,
                    image_cdn,
                    timeout,
                    dict(session.headers),
                ): slot
                for slot in batch
            }
            for fut in as_completed(futures):
                slot = futures[fut]
                name = slot["name"]
                try:
                    results[name] = fut.result()
                except Exception as e:
                    results[name] = {"ok": False, "attempts": [], "error": str(e)}
                checked += 1
                elapsed = time.time() - img_started
                remaining = max(0, total_slots - checked)
                avg = elapsed / max(1, checked - already_processed)
                eta = avg * remaining
                outcome = "HIT" if results[name].get("ok") else "MISS"
                log(
                    f"    [IMG {checked}/{total_slots} | saved={len(manifest['downloaded'])}/{target_label}] "
                    f"{name} -> {outcome}; elapsed={format_duration(elapsed)}; ETA~{format_duration(eta)}"
                )

        # Commit in CL rank order even though network requests finished out of order.
        for slot in batch:
            name = slot["name"]
            result = results.get(name, {"ok": False, "attempts": []})
            processed_names.add(name)
            if not result.get("ok"):
                manifest["failed_slots"].append({
                    "name": name,
                    "rank": slot["rank"],
                    "attempts": result.get("attempts", []),
                })
                continue

            content = result["content"]
            h = hashlib.sha256(content).hexdigest()
            cand = result.get("candidate") or {}
            if h in seen_hashes:
                manifest["duplicate_content"].append({
                    "name": name,
                    "path": cand.get("path"),
                    "url": result.get("url"),
                    "sha256": h,
                })
                continue

            if target is not None and len(manifest["downloaded"]) >= target:
                # The request was already in-flight in the final concurrent batch.
                # Do not persist more than the requested target.
                continue

            seen_hashes.add(h)
            idx = len(manifest["downloaded"]) + 1
            ext = ext_from_content_type(result.get("content_type", ""), result.get("url", ""))
            dest = out_dir / f"{idx:02d}_{safe_name(name)}{ext}"
            dest.write_bytes(content)
            item = {
                "index": idx,
                "name": name,
                "rank": slot["rank"],
                "logical_path": cand.get("path"),
                "localized": bool(cand.get("localized")),
                "url": result.get("url"),
                "file": dest.name,
                "sha256": h,
            }
            manifest["downloaded"].append(item)
            log(
                f"      -> SAVED {len(manifest['downloaded'])}/{target_label}: "
                f"{dest.name} ({'localized' if cand.get('localized') else 'generic'})"
            )

        manifest["processed_slots"] = sorted(processed_names)
        manifest["arkswift_upload_candidates"] = [x["file"] for x in manifest["downloaded"][:12]]
        _write_manifest(manifest_path, manifest)  # checkpoint after every concurrent batch

    manifest["complete"] = True
    manifest["processed_slots"] = sorted(processed_names)
    manifest["arkswift_upload_candidates"] = [x["file"] for x in manifest["downloaded"][:12]]
    normalized["image_source"]["downloaded_files"] = [x["file"] for x in manifest["downloaded"]]
    normalized["image_source"]["arkswift_upload_candidates"] = manifest["arkswift_upload_candidates"]

    if download_gpsr and info["gpsr"]:
        gpsr_dir = sku_dir / "gpsr"
        gpsr_dir.mkdir(exist_ok=True)
        manifest["gpsr"] = []
        gpsr_seen = set()
        for j, row in enumerate(info["gpsr"], 1):
            secret = clean_text(row.get("PIC_URL_SECRET"))
            if not secret:
                continue
            url = image_cdn.rstrip("/") + "/" + secret.lstrip("/")
            dest = download_image(session, url, gpsr_dir / f"gpsr_{j:02d}", timeout=timeout)
            if not dest:
                continue
            h = hashlib.sha256(dest.read_bytes()).hexdigest()
            if h in gpsr_seen:
                dest.unlink(missing_ok=True)
                continue
            gpsr_seen.add(h)
            manifest["gpsr"].append({"path": row.get("PIC_URL"), "url": url, "file": str(Path("gpsr") / dest.name)})

    _write_manifest(manifest_path, manifest)
    log(
        f"  [IMAGES DONE] saved={len(manifest['downloaded'])}/{target_label}; "
        f"checked={len(processed_names)}/{total_slots}; elapsed={format_duration(time.time()-img_started)}"
    )
    return manifest

def collect_legacy_images(
    session: requests.Session,
    normalized: Dict[str, Any],
    out_dir: Path,
    max_images: int,
    timeout: int,
) -> Dict[str, Any]:
    """Backward-compatible GetDetail direct-image mode."""
    out_dir.mkdir(parents=True, exist_ok=True)
    manifest: Dict[str, Any] = {"mode": "main", "downloaded": [], "failed_direct": []}
    seen_hashes = set()
    index = 1
    for url in normalized["image_source"]["direct_urls_from_getdetail"]:
        if max_images > 0 and len(manifest["downloaded"]) >= max_images:
            break
        dest = download_image(session, url, out_dir / f"{index:02d}_getdetail", timeout=timeout)
        if dest:
            h = hashlib.sha256(dest.read_bytes()).hexdigest()
            if h in seen_hashes:
                dest.unlink(missing_ok=True)
                continue
            seen_hashes.add(h)
            manifest["downloaded"].append({"kind": "getdetail", "url": url, "file": dest.name, "sha256": h})
            index += 1
        else:
            manifest["failed_direct"].append(url)
    (out_dir / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8")
    return manifest


# ----------------------------
# Summary CSV
# ----------------------------

SUMMARY_FIELDS = [
    "requested_sku", "returned_sku", "status",
    "product_title", "product_title_original", "item_name_en", "item_name_cn", "brand",
    "category_1", "category_2", "category_3", "product_type", "product_type_name",
    "arkswift_category", "seller_sku", "ean_upc", "country_of_origin",
    "color", "material", "is_white_label", "is_original",
    "length_cm", "width_cm", "height_cm", "net_weight_kg",
    "box_qty", "box_length_cm", "box_width_cm", "box_height_cm", "gross_weight_kg",
    "feature_1", "feature_2", "feature_3", "feature_4", "feature_5",
    "description", "specifications", "package_includes", "electric_flag", "manual_url",
    "images_downloaded", "images_dir", "image_files", "missing_required", "warnings", "error",
]


def summary_row(n: Optional[Dict[str, Any]], status: str, error: str = "", images: int = 0) -> Dict[str, Any]:
    row = {k: "" for k in SUMMARY_FIELDS}
    row["status"] = status
    row["error"] = error
    row["images_downloaded"] = images
    if not n:
        return row

    ensure_brand_free_title(n)
    a = n["arkswift_fields"]
    feats = list(a.get("features") or [])[:5]
    feats += [""] * (5 - len(feats))
    sku = n["source"]["requested_sku"]
    image_files = list((n.get("image_source") or {}).get("arkswift_upload_candidates") or [])
    row.update({
        "requested_sku": sku,
        "returned_sku": n["identity"].get("seller_sku"),
        "product_title": a.get("product_title"),
        "product_title_original": a.get("product_title_original"),
        "item_name_en": n["identity"].get("item_name_en"),
        "item_name_cn": n["identity"].get("item_name_cn"),
        "brand": n["identity"].get("brand"),
        "category_1": n["category_source"].get("level1"),
        "category_2": n["category_source"].get("level2"),
        "category_3": n["category_source"].get("level3"),
        "product_type": n["category_source"].get("product_type"),
        "product_type_name": n["category_source"].get("product_type_name"),
        # Final ArkSwift category mapping is a later rule layer; keep an explicit column.
        "arkswift_category": a.get("arkswift_category") or "",
        "seller_sku": a.get("seller_sku"),
        "ean_upc": "|".join(str(x) for x in (a.get("upc_ean_candidates") or [])),
        "country_of_origin": a.get("country_of_origin"),
        "color": a.get("color_source"),
        "material": a.get("material_source"),
        "is_white_label": "" if a.get("is_white_label") is None else a.get("is_white_label"),
        "is_original": "" if a.get("is_original") is None else a.get("is_original"),
        "length_cm": (a.get("assembled_dimensions_cm") or {}).get("length"),
        "width_cm": (a.get("assembled_dimensions_cm") or {}).get("width"),
        "height_cm": (a.get("assembled_dimensions_cm") or {}).get("height"),
        "net_weight_kg": a.get("net_weight_kg"),
        "box_qty": (a.get("package") or {}).get("quantity"),
        "box_length_cm": (a.get("package") or {}).get("length_cm"),
        "box_width_cm": (a.get("package") or {}).get("width_cm"),
        "box_height_cm": (a.get("package") or {}).get("height_cm"),
        "gross_weight_kg": (a.get("package") or {}).get("gross_weight_kg"),
        "feature_1": feats[0], "feature_2": feats[1], "feature_3": feats[2],
        "feature_4": feats[3], "feature_5": feats[4],
        "description": a.get("description"),
        "specifications": a.get("specifications"),
        "package_includes": a.get("package_includes"),
        "electric_flag": a.get("electric_flag"),
        "manual_url": a.get("manual_url"),
        "images_dir": str(Path(safe_name(sku)) / "images"),
        "image_files": "|".join(image_files),
        "missing_required": "|".join((n.get("validation") or {}).get("missing_required") or []),
        "warnings": "|".join((n.get("validation") or {}).get("warnings") or []),
    })
    return row


def _load_json(path: Path) -> Optional[Any]:
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return None


def existing_sku_artifacts(sku_dir: Path, images_mode: str) -> Tuple[Optional[Dict[str, Any]], Optional[Dict[str, Any]], bool]:
    """Return (product, manifest, fully_done) for resume/skip logic."""
    product_path = sku_dir / "product.json"
    product = _load_json(product_path) if product_path.exists() else None
    if not isinstance(product, dict):
        product = None

    if images_mode == "none":
        return product, None, bool(product)

    manifest_path = sku_dir / "images" / "manifest.json"
    manifest = _load_json(manifest_path) if manifest_path.exists() else None
    if not isinstance(manifest, dict):
        manifest = None

    # Old versions wrote manifest.json only after finishing, so absence of `complete`
    # means a pre-0.5 completed image stage. New 0.5 runs checkpoint with complete=false.
    if manifest is not None:
        complete = manifest.get("complete")
        if complete is None:
            complete = True
        return product, manifest, bool(product and complete)
    return product, None, False


# ----------------------------
# Main
# ----------------------------

def main() -> int:
    ap = argparse.ArgumentParser(description="RevFlow")
    ap.add_argument("--force", action="store_true")
    cli = ap.parse_args()

    cfg = load_config()
    country_prefix = cfg["market"]
    country = country_prefix.upper()
    market_config(country_prefix)
    paths = workspace_paths(country_prefix)
    detect_manifest = validate_detect(country_prefix)
    input_path = paths["detect"] / f"{country_prefix}-all.txt"
    todo_path = paths["detect"] / f"{country_prefix}-need_create.txt"
    skus = load_skus(input_path)
    todo_set = {x.upper() for x in load_skus(todo_path)}
    if not skus:
        raise ValueError("SKU Detect 的完整清单为空，请重新检测。")
    if todo_set - {x.upper() for x in skus}:
        raise ValueError("need_create 含有完整清单以外的 SKU，请重新检测。")

    args = SimpleNamespace(
        curl=PROJECT_ROOT / "revflow" / "getdetail.curl.txt",
        input=input_path,
        sheet=None,

        out=paths["extract"],
        country=country,

        delay=float(cfg.get("delay", 0.8)),
        timeout=int(cfg.get("timeout", 30)),
        image_timeout=int(cfg.get("image_timeout", 20)),

        images="exact",
        max_images=int(cfg.get("max_images", 12)),
        image_workers=int(cfg.get("image_workers", 4)),

        force=cli.force,
        image_cdn=DEFAULT_IMAGE_CDN,
        download_gpsr=False,
    )

    checkpoint_every = max(
        1,
        int(cfg.get("checkpoint_every", 10))
    )
    if args.max_images == 0:
        args.max_images = 12

    args.out.mkdir(parents=True, exist_ok=True)

    write_extract_context(country_prefix, detect_manifest, complete=False)
    if not args.curl.is_file():
        raise FileNotFoundError("缺少 revflow/getdetail.curl.txt；请在 CL 对应国家复制 GetDetail 请求为 cURL。")
    seed_url, headers = parse_curl_file(args.curl)
    if "getdetail" not in seed_url.lower():
        log("[WARN] Seed URL does not contain 'GetDetail'. Make sure you copied the correct XHR.")
    session = make_session(headers, cookie_domain=urlsplit(seed_url).hostname)

    log(f"[INFO] input worklist: {input_path}")
    log(f"[INFO] sku_detect mask: {todo_path} | to_do={len(todo_set)} | skip_existing={len(skus) - len(todo_set)}")

    effective_target = "ALL" if args.max_images < 0 else (12 if args.max_images == 0 else args.max_images)
    log(f"[INFO] SKUs: {len(skus)} | country={args.country} | images={args.images} | image_target={effective_target}")
    log(f"[INFO] CL API mode: SERIAL / low-frequency | image CDN workers={max(1, min(args.image_workers, 8))}")
    log("[INFO] resume: completed output is skipped; product-only output resumes at image stage")
    log("[INFO] lineGuid resolver: CountryList/GetPageList (automatic for non-seed SKUs)")
    log(f"[INFO] Output: {args.out.resolve()}")

    summary: List[Dict[str, Any]] = []
    summary_path = args.out / "summary.csv"
    batch_started = time.time()

    for i, sku in enumerate(skus, 1):
        sku = sku.strip()
        sku_dir = args.out / safe_name(sku)
        cl_work_attempted = False
        log(f"[{i}/{len(skus)}] {sku}")

        try:
            # Fast gate: sku_detect says this SKU already has an ArkSwift review/product
            # record in the CURRENT market. Keep it in summary.csv (original order),
            # but do not touch CL and do not delete any old cached product data.
            if sku.upper() not in todo_set:
                row = summary_row(None, status="SKIP_EXISTS")
                row["requested_sku"] = sku
                row["seller_sku"] = sku
                row["warnings"] = "sku_detect: existing in target ArkSwift store; CL skipped"
                summary.append(row)
                log("  [SKIP_EXISTS] sku_detect 已确认目标店铺存在记录；跳过 CL")
                raise StopIteration

            sku_dir.mkdir(parents=True, exist_ok=True)
            existing_product, existing_manifest, fully_done = existing_sku_artifacts(sku_dir, args.images)
            if fully_done and not args.force:
                ensure_brand_free_title(existing_product)
                image_count = len((existing_manifest or {}).get("downloaded", []))
                if existing_manifest:
                    image_source = existing_product.setdefault("image_source", {})
                    candidates = list(existing_manifest.get("arkswift_upload_candidates") or [])
                    if not candidates:
                        candidates = [x.get("file") for x in existing_manifest.get("downloaded", [])[:12] if x.get("file")]
                    image_source["arkswift_upload_candidates"] = candidates[:12]
                    image_source["downloaded_files"] = [
                        x.get("file") for x in existing_manifest.get("downloaded", []) if x.get("file")
                    ]
                # Persist migration of old product.json title + image candidate metadata.
                (sku_dir / "product.json").write_text(
                    json.dumps(existing_product, ensure_ascii=False, indent=2), encoding="utf-8"
                )
                # Cached CL output is still actionable downstream. Mark it OK rather
                # than SKIPPED so sku_write can reuse it without re-querying CL.
                summary.append(summary_row(existing_product, status="OK", images=image_count))
                log(f"  [CACHE] completed CL output reused | images={image_count}")
                raise StopIteration

            seed_sku = clean_text(_query_value(seed_url, "keyValue"))
            resolved_line_guid = None

            if existing_product and not args.force:
                normalized = existing_product
                ensure_brand_free_title(normalized)
                log("  [RESUME] product.json exists; skipping GetPageList/GetDetail and resuming image stage")
            elif seed_sku and seed_sku.upper() == sku.upper():
                cl_work_attempted = True
                resolved_line_guid = clean_text(_query_value(seed_url, "lineGuid"))
                log(f"  [LINEGUID] seed SKU -> {resolved_line_guid or 'missing'}")
            else:
                cl_work_attempted = True
                log("  [LINEGUID] resolving via CountryList/GetPageList ...")
                try:
                    resolved_line_guid, page_payload, line_debug = resolve_line_guid_via_getpagelist(
                        session=session,
                        seed_url=seed_url,
                        sku=sku,
                        timeout=args.timeout,
                    )
                    (sku_dir / "raw_getpagelist.json").write_text(
                        json.dumps(page_payload, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
                    (sku_dir / "lineguid_resolution.json").write_text(
                        json.dumps(line_debug, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
                    log(f"  [LINEGUID] OK -> {resolved_line_guid}")
                except LineGuidResolutionError as le:
                    if le.payload is not None:
                        (sku_dir / "raw_getpagelist.json").write_text(
                            json.dumps(le.payload, ensure_ascii=False, indent=2), encoding="utf-8"
                        )
                    (sku_dir / "lineguid_resolution.json").write_text(
                        json.dumps(le.debug, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
                    raise

            if not (existing_product and not args.force):
                payload, detail_url, detail_attempts = fetch_getdetail_for_sku(
                    session=session,
                    seed_url=seed_url,
                    sku=sku,
                    timeout=args.timeout,
                    resolved_line_guid=resolved_line_guid,
                )
                (sku_dir / "getdetail_attempts.json").write_text(
                    json.dumps(detail_attempts, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
                (sku_dir / "raw_getdetail.json").write_text(
                    json.dumps(payload, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )

                normalized = normalize_product(payload, sku, args.country)

                returned = (normalized["identity"]["seller_sku"] or "").upper()
                if returned != sku.upper():
                    raise RuntimeError(
                        f"SKU validation failed: requested={sku}, returned={normalized['identity']['seller_sku']}"
                    )

                # Keep the resolved master line GUID for audit/debugging.
                normalized["source"]["resolved_master_line_guid"] = clean_text(
                    find_master(payload).get("MASTER_LINE_GUID")
                )
                normalized["source"]["line_guid_resolution"] = (
                    "seed-curl" if seed_sku and seed_sku.upper() == sku.upper() else "CountryList/GetPageList"
                )

            # Persist the structured CL data immediately. Image extraction is a second stage
            # and must not erase a successful product-data extraction.
            (sku_dir / "product.json").write_text(
                json.dumps(normalized, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )

            img_manifest = {"downloaded": []}
            image_error = ""
            try:
                if args.images == "exact":
                    if not fully_done:
                        cl_work_attempted = True
                    img_manifest = collect_exact_images(
                        session=session,
                        seed_url=seed_url,
                        normalized=normalized,
                        sku_dir=sku_dir,
                        max_images=args.max_images,
                        image_cdn=args.image_cdn,
                        download_gpsr=args.download_gpsr,
                        timeout=args.image_timeout,
                        image_workers=args.image_workers,
                    )
                elif args.images == "main":
                    img_manifest = collect_legacy_images(
                        session=session,
                        normalized=normalized,
                        out_dir=sku_dir / "images",
                        max_images=args.max_images,
                        timeout=args.image_timeout,
                    )
            except Exception as ie:
                image_error = str(ie)
                log(f"  [IMAGE ERROR] {image_error}")

            # Write again after image discovery so logical_photo_list/candidates are included.
            (sku_dir / "product.json").write_text(
                json.dumps(normalized, ensure_ascii=False, indent=2),
                encoding="utf-8",
            )

            if image_error:
                summary.append(summary_row(
                    normalized,
                    status="PARTIAL",
                    error=image_error,
                    images=len(img_manifest.get("downloaded", [])),
                ))
            else:
                summary.append(summary_row(
                    normalized,
                    status="OK",
                    images=len(img_manifest.get("downloaded", [])),
                ))

        except StopIteration:
            pass
        except Exception as e:
            log(f"  [ERROR] {e}")
            if isinstance(e, GetDetailResolutionError):
                (sku_dir / "getdetail_attempts.json").write_text(
                    json.dumps(e.attempts, ensure_ascii=False, indent=2),
                    encoding="utf-8",
                )
            if isinstance(e, LineGuidResolutionError):
                if e.payload is not None:
                    (sku_dir / "raw_getpagelist.json").write_text(
                        json.dumps(e.payload, ensure_ascii=False, indent=2), encoding="utf-8"
                    )
                (sku_dir / "lineguid_resolution.json").write_text(
                    json.dumps(e.debug, ensure_ascii=False, indent=2), encoding="utf-8"
                )
            summary.append(summary_row(None, status="ERROR", error=str(e)))
            summary[-1]["requested_sku"] = sku

        elapsed_batch = time.time() - batch_started
        avg_sku = elapsed_batch / i
        eta_batch = avg_sku * max(0, len(skus) - i)
        log(
            f"  [BATCH] completed={i}/{len(skus)} | "
            f"elapsed(total)={format_duration(elapsed_batch)} | "
            f"avg/SKU={format_duration(avg_sku)} | "
            f"ETA(total)~{format_duration(eta_batch)}"
        )
        if i % checkpoint_every == 0:
            with summary_path.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.DictWriter(f, fieldnames=SUMMARY_FIELDS)
                w.writeheader()
                w.writerows(summary)
            log(f"  [CHECKPOINT] summary.csv saved | rows={len(summary)}")

        # No artificial delay for rows that were only tagged SKIP_EXISTS or reused
        # from a fully completed local cache. Delay remains for real CL/image work.
        if i < len(skus) and cl_work_attempted:
            time.sleep(max(0, args.delay))


    with summary_path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=SUMMARY_FIELDS)
        w.writeheader()
        w.writerows(summary)

    ok = sum(1 for x in summary if x["status"] == "OK")
    partial = sum(1 for x in summary if x["status"] == "PARTIAL")
    skipped = sum(1 for x in summary if x["status"] in {"SKIPPED", "SKIP_EXISTS"})
    failed = sum(1 for x in summary if x["status"] == "ERROR")
    log(f"[DONE] OK={ok} SKIP_EXISTS={skipped} PARTIAL={partial} ERROR={failed} / {len(summary)}")
    log(f"[DONE] Summary: {summary_path.resolve()}")
    write_extract_context(country_prefix, detect_manifest, complete=True)
    return 0 if failed == 0 and partial == 0 else 1


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\n[STOP] 已中断；修复问题后重跑相同命令，复用已完成的商品/图片缓存。", file=sys.stderr)
        raise SystemExit(130)
    except Exception as exc:
        print(f"[ERROR] {exc}", file=sys.stderr)
        raise SystemExit(1)
